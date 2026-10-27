"""Connector World Bank Commodity Markets (Pink Sheet): giá dầu, vàng quốc tế theo tháng.

Đường dẫn file chứa mã đổi theo kỳ phát hành, nên tìm link trên trang thay vì ghi cứng.
"""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from macro.connectors.base import Observation, SourceError
from macro.connectors.dates import month_end
from macro.connectors.http import download, get
from macro.series import Series, SeriesStatus

PAGE_URL = "https://www.worldbank.org/en/research/commodity-markets"
FILENAME = "CMO-Historical-Data-Monthly.xlsx"
LINK = re.compile(r"https?://[^\"'\s>]*CMO-Historical-Data-Monthly\.xlsx")
SHEET = "Monthly Prices"
PERIOD = re.compile(r"^(\d{4})M(\d{2})$")


def _series(id: str, column: str, name: str, unit: str, order: int) -> Series:
    return Series(
        id=id,
        name=name,
        measures=f"Pink Sheet, cột '{column}', giá danh nghĩa trung bình tháng",
        connector="wb_pinksheet",
        upstream_provider="World Bank Commodity Markets",
        source_series_id=column,
        frequency="M",
        unit=unit,
        transform="level",
        status=SeriesStatus.VERIFIED,
        source_url=PAGE_URL,
        terms_note=(
            "World Bank: dữ liệu công khai; giá quốc tế theo tháng, không phải giá trong nước."
        ),
        display_order=order,
    )


SERIES = [
    _series("global.brent", "Crude oil, Brent", "Dầu Brent", "USD/thùng", 210),
    _series("global.gold", "Gold", "Vàng", "USD/troy ounce", 220),
]


class PinkSheetConnector:
    name = "wb_pinksheet"

    def series(self) -> list[Series]:
        return SERIES

    def fetch(self, raw_dir: Path) -> None:
        match = LINK.search(get(PAGE_URL).text)
        if match is None:
            raise SourceError(f"không tìm thấy link {FILENAME} trên {PAGE_URL}")
        download(match.group(0), raw_dir / FILENAME)

    def parse(self, raw_dir: Path) -> list[Observation]:
        path = raw_dir / FILENAME
        if not path.is_file():
            raise SourceError(f"thiếu {FILENAME}")
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            raise SourceError(f"{FILENAME} không đọc được: {exc}") from exc
        if SHEET not in workbook.sheetnames:
            raise SourceError(f"{FILENAME}: không có sheet '{SHEET}'")
        rows = list(workbook[SHEET].iter_rows(values_only=True))
        header = next((row for row in rows if "Crude oil, Brent" in row), None)
        if header is None:
            raise SourceError(f"{FILENAME}: không thấy dòng tên cột")
        columns = {}
        for item in SERIES:
            if item.source_series_id not in header:
                raise SourceError(f"{FILENAME}: không có cột '{item.source_series_id}'")
            columns[item.id] = header.index(item.source_series_id)
        observations = []
        for row in rows:
            match = PERIOD.match(str(row[0] or ""))
            if match is None:
                continue
            period = month_end(int(match.group(1)), int(match.group(2)))
            for series_id, index in columns.items():
                value = row[index]
                if isinstance(value, int | float):
                    observations.append(Observation(series_id, period, float(value)))
        return observations
